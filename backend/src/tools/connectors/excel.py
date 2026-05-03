"""
FACT System Excel Connector Tool

This module implements the Excel connector for secure workbook inspection
and sheet preview, following the same tool structure as sql.py.

MVP scope:
- Open .xlsx files
- List sheets
- Infer columns from the header row
- Preview rows from a selected sheet

Not included yet:
- Upload endpoint
- Persistent per-user workbook storage
- NL-to-formula / NL-to-sheet-query translation
"""

import os
import time
from typing import Dict, Any, List, Optional
import structlog

try:
    # Try relative imports first (when used as package)
    from ...core.errors import DatabaseError, SecurityError
    from ..decorators import Tool
except ImportError:
    # Fall back to absolute imports (when run as script)
    import sys
    from pathlib import Path

    src_path = str(Path(__file__).parent.parent.parent)
    if src_path not in sys.path:
        sys.path.insert(0, src_path)

    from core.errors import DatabaseError, SecurityError
    from tools.decorators import Tool

try:
    from openpyxl import load_workbook
except ImportError as e:
    raise ImportError(
        "openpyxl is required for the Excel connector. Run: pip install openpyxl"
    ) from e


logger = structlog.get_logger(__name__)


class ExcelTool:
    """
    Excel connector tool for secure workbook inspection.

    Supports:
    - workbook overview
    - sheet preview
    - column inference from first row
    """

    def __init__(self) -> None:
        self._metadata_cache: Dict[str, Dict[str, Any]] = {}
        self._cache_max_size = 200

    # ------------------------------------------------------------------
    # Internal validation helpers
    # ------------------------------------------------------------------

    def _validate_file_path(self, file_path: str) -> None:
        """
        Validate the workbook path for basic security.

        Args:
            file_path: Path to .xlsx file

        Raises:
            SecurityError: If path is unsafe
            DatabaseError: If file is missing / invalid
        """
        if not isinstance(file_path, str) or not file_path.strip():
            raise DatabaseError("Excel file path is required")

        normalized = os.path.abspath(file_path)

        if not os.path.exists(normalized):
            raise DatabaseError("Excel file does not exist")

        if not os.path.isfile(normalized):
            raise DatabaseError("Provided path is not a file")

        if not normalized.lower().endswith(".xlsx"):
            raise SecurityError("Only .xlsx files are supported")

        # Basic size check for MVP safety
        size_bytes = os.path.getsize(normalized)
        max_size_bytes = 25 * 1024 * 1024  # 25 MB
        if size_bytes > max_size_bytes:
            raise SecurityError("Excel file is too large. Max supported size is 25 MB")

    def _validate_sheet_name(self, sheet_name: str) -> None:
        """
        Validate sheet name.

        Args:
            sheet_name: Workbook sheet name

        Raises:
            SecurityError: If invalid
        """
        if not isinstance(sheet_name, str) or not sheet_name.strip():
            raise SecurityError("Sheet name is required")

        if len(sheet_name) > 128:
            raise SecurityError("Sheet name is too long")

    def _normalize_cell_value(self, value: Any) -> Any:
        """
        Convert workbook cell values into JSON-friendly values.
        """
        if value is None:
            return None

        # Excel formulas are already resolved because we load with data_only=True
        if isinstance(value, (str, int, float, bool)):
            return value

        # Dates, decimals, etc. -> string fallback for MVP
        return str(value)

    def _safe_header_name(self, value: Any, index: int) -> str:
        """
        Turn header cell into a usable column name.
        """
        if value is None:
            return f"column_{index + 1}"

        text = str(value).strip()
        if not text:
            return f"column_{index + 1}"

        return text

    def _get_cached_metadata(self, file_path: str) -> Optional[Dict[str, Any]]:
        return self._metadata_cache.get(file_path)

    def _set_cached_metadata(self, file_path: str, metadata: Dict[str, Any]) -> None:
        if len(self._metadata_cache) >= self._cache_max_size:
            oldest_keys = list(self._metadata_cache.keys())[:20]
            for key in oldest_keys:
                del self._metadata_cache[key]
        self._metadata_cache[file_path] = metadata

    # ------------------------------------------------------------------
    # Workbook operations
    # ------------------------------------------------------------------

    def _load_workbook(self, file_path: str):
        """
        Open workbook in read-only mode.
        """
        self._validate_file_path(file_path)
        try:
            return load_workbook(
                filename=file_path,
                read_only=True,
                data_only=True
            )
        except Exception as e:
            raise DatabaseError(f"Failed to open Excel workbook: {e}")

    async def get_workbook_overview(self, file_path: str) -> Dict[str, Any]:
        """
        Return workbook-level metadata and sheet summaries.
        """
        start_time = time.time()
        self._validate_file_path(file_path)

        cached = self._get_cached_metadata(file_path)
        if cached is not None:
            return {
                **cached,
                "cached": True,
                "execution_time_ms": round((time.time() - start_time) * 1000, 2),
                "status": "success",
            }

        workbook = None
        try:
            workbook = self._load_workbook(file_path)
            sheets_summary: List[Dict[str, Any]] = []

            for ws in workbook.worksheets:
                max_rows = ws.max_row or 0
                max_cols = ws.max_column or 0

                header_row = []
                if max_rows >= 1:
                    first_row = next(
                        ws.iter_rows(min_row=1, max_row=1, values_only=True),
                        tuple()
                    )
                    header_row = [
                        self._safe_header_name(value, idx)
                        for idx, value in enumerate(first_row)
                    ]

                sheets_summary.append({
                    "sheet_name": ws.title,
                    "row_count_estimate": max_rows,
                    "column_count_estimate": max_cols,
                    "columns": header_row,
                })

            response = {
                "file_path": file_path,
                "file_name": os.path.basename(file_path),
                "sheet_count": len(workbook.sheetnames),
                "sheets": sheets_summary,
            }

            self._set_cached_metadata(file_path, response)

            return {
                **response,
                "cached": False,
                "execution_time_ms": round((time.time() - start_time) * 1000, 2),
                "status": "success",
            }

        except (SecurityError, DatabaseError):
            raise
        except Exception as e:
            raise DatabaseError(f"Failed to inspect workbook: {e}")
        finally:
            if workbook is not None:
                workbook.close()

    async def get_sheet_preview(
        self,
        file_path: str,
        sheet_name: str,
        max_rows: int = 20
    ) -> Dict[str, Any]:
        """
        Return a structured preview of one sheet.

        Args:
            file_path: Path to workbook
            sheet_name: Name of sheet to preview
            max_rows: Max number of data rows to return (excluding header)
        """
        start_time = time.time()

        self._validate_file_path(file_path)
        self._validate_sheet_name(sheet_name)

        if not isinstance(max_rows, int) or max_rows < 1 or max_rows > 100:
            raise SecurityError("max_rows must be between 1 and 100")

        workbook = None
        try:
            workbook = self._load_workbook(file_path)

            if sheet_name not in workbook.sheetnames:
                raise DatabaseError(f"Sheet '{sheet_name}' does not exist")

            ws = workbook[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)

            header_tuple = next(rows_iter, tuple())
            columns = [
                self._safe_header_name(value, idx)
                for idx, value in enumerate(header_tuple)
            ]

            preview_rows: List[Dict[str, Any]] = []
            count = 0

            for row in rows_iter:
                if count >= max_rows:
                    break

                row_dict = {}
                for idx, col_name in enumerate(columns):
                    value = row[idx] if idx < len(row) else None
                    row_dict[col_name] = self._normalize_cell_value(value)

                # Skip completely empty rows
                if any(v is not None and str(v).strip() != "" for v in row_dict.values()):
                    preview_rows.append(row_dict)
                    count += 1

            return {
                "file_path": file_path,
                "file_name": os.path.basename(file_path),
                "sheet_name": sheet_name,
                "columns": columns,
                "rows": preview_rows,
                "row_count": len(preview_rows),
                "execution_time_ms": round((time.time() - start_time) * 1000, 2),
                "status": "success",
            }

        except (SecurityError, DatabaseError):
            raise
        except Exception as e:
            raise DatabaseError(f"Failed to preview sheet: {e}")
        finally:
            if workbook is not None:
                workbook.close()


# ---------------------------------------------------------------------------
# Error sanitization
# ---------------------------------------------------------------------------

def _sanitize_excel_error(exc: Exception) -> str:
    """
    Convert raw Excel-related exceptions into short, LLM-friendly messages.
    """
    raw = str(exc).strip()

    if "does not exist" in raw.lower():
        return "Error: Excel file or sheet does not exist."[:100]

    if "only .xlsx files are supported" in raw.lower():
        return "Error: only .xlsx files are supported."[:100]

    if "too large" in raw.lower():
        return "Error: Excel file is too large."[:100]

    if "failed to open" in raw.lower():
        return "Error: failed to open Excel workbook."[:100]

    if "sheet" in raw.lower() and "does not exist" in raw.lower():
        return "Error: requested sheet does not exist."[:100]

    first_line = raw.splitlines()[0] if raw else "unknown error"
    generic = f"Error: {first_line}" if first_line else "Error: Excel request failed."
    return generic[:100]


# Global Excel tool instance
_excel_tool_instance: Optional[ExcelTool] = None


def initialize_excel_tool() -> None:
    """
    Initialize the global Excel tool instance.
    """
    global _excel_tool_instance
    _excel_tool_instance = ExcelTool()
    logger.info("Excel tool initialized")


def get_excel_tool() -> ExcelTool:
    """
    Return the initialized global Excel tool instance.
    """
    if _excel_tool_instance is None:
        raise RuntimeError("Excel tool not initialized. Call initialize_excel_tool first.")
    return _excel_tool_instance


# ---------------------------------------------------------------------------
# Tool definitions
# ---------------------------------------------------------------------------

@Tool(
    name="Excel_GetWorkbookOverview",
    description=(
        "Inspect an Excel .xlsx workbook and return workbook overview, "
        "sheet names, estimated row/column counts, and inferred headers."
    ),
    parameters={
        "file_path": {
            "type": "string",
            "description": "Absolute or server-accessible path to the .xlsx file to inspect.",
            "minLength": 5,
            "maxLength": 1000,
        }
    },
    requires_auth=False,
    timeout_seconds=20,
    version="1.0.0"
)
async def excel_get_workbook_overview(file_path: str) -> Dict[str, Any]:
    """
    Inspect workbook and return high-level metadata.

    Example:
        result = await excel_get_workbook_overview("/tmp/sales.xlsx")
    """
    if _excel_tool_instance is None:
        return {
            "error": "Excel tool not initialized",
            "status": "failed",
            "execution_time_ms": 0,
        }

    try:
        return await _excel_tool_instance.get_workbook_overview(file_path)
    except (SecurityError, DatabaseError) as e:
        return {
            "error": _sanitize_excel_error(e),
            "status": "failed",
            "execution_time_ms": 0,
        }
    except Exception as e:
        logger.error("Unexpected error during workbook overview", error=str(e))
        return {
            "error": _sanitize_excel_error(e),
            "status": "failed",
            "execution_time_ms": 0,
        }


@Tool(
    name="Excel_GetSheetPreview",
    description=(
        "Preview rows from a specific sheet in an Excel .xlsx workbook. "
        "Uses the first row as headers and returns structured row objects."
    ),
    parameters={
        "file_path": {
            "type": "string",
            "description": "Absolute or server-accessible path to the .xlsx file.",
            "minLength": 5,
            "maxLength": 1000,
        },
        "sheet_name": {
            "type": "string",
            "description": "Name of the sheet to preview.",
            "minLength": 1,
            "maxLength": 128,
        },
        "max_rows": {
            "type": "integer",
            "description": "Maximum number of data rows to preview, between 1 and 100.",
            "minimum": 1,
            "maximum": 100,
        },
    },
    requires_auth=False,
    timeout_seconds=20,
    version="1.0.0"
)
async def excel_get_sheet_preview(
    file_path: str,
    sheet_name: str,
    max_rows: int = 20
) -> Dict[str, Any]:
    """
    Preview rows from a selected sheet.

    Example:
        result = await excel_get_sheet_preview("/tmp/sales.xlsx", "Sheet1", 10)
    """
    if _excel_tool_instance is None:
        return {
            "error": "Excel tool not initialized",
            "status": "failed",
            "execution_time_ms": 0,
        }

    try:
        return await _excel_tool_instance.get_sheet_preview(
            file_path=file_path,
            sheet_name=sheet_name,
            max_rows=max_rows,
        )
    except (SecurityError, DatabaseError) as e:
        return {
            "error": _sanitize_excel_error(e),
            "status": "failed",
            "execution_time_ms": 0,
        }
    except Exception as e:
        logger.error("Unexpected error during sheet preview", error=str(e))
        return {
            "error": _sanitize_excel_error(e),
            "status": "failed",
            "execution_time_ms": 0,
        }


@Tool(
    name="Excel_GetUsageGuide",
    description="Get sample usage guidance for the Excel connector tools.",
    parameters={},
    requires_auth=False,
    timeout_seconds=5,
    version="1.0.0"
)
async def excel_get_usage_guide() -> Dict[str, Any]:
    """
    Return examples showing how the Excel connector should be used.
    """
    return {
        "examples": [
            {
                "step": 1,
                "tool": "Excel_GetWorkbookOverview",
                "purpose": "List workbook sheets and detect available columns",
                "example_call": {
                    "file_path": "/tmp/financial_model.xlsx"
                }
            },
            {
                "step": 2,
                "tool": "Excel_GetSheetPreview",
                "purpose": "Preview rows from one sheet before deeper analysis",
                "example_call": {
                    "file_path": "/tmp/financial_model.xlsx",
                    "sheet_name": "Revenue",
                    "max_rows": 10
                }
            }
        ],
        "status": "success",
        "execution_time_ms": 0
    }